"""Construtor de planilhas sinteticas para os testes de integracao.

Nenhuma planilha real de cliente entra no repositorio: os arquivos sao gerados
em tmp_path a cada teste. Todos os CNPJs, nomes e certificados sao ficticios.
"""
import openpyxl

CABECALHO_EMPRESAS = ["CNPJ", "EMPRESA", "CERTIFICADO", "DCTFWEB", "PROCESSOS"]

# Sentinelas ficticias — nenhuma corresponde a empresa ou documento real.
ALFA = ("11111111000191", "ALFA FICTICIA LTDA", "CERT ALFA")
BETA = ("22222222000172", "BETA FICTICIA LTDA", "CERT BETA")
GAMA = ("33333333000153", "GAMA FICTICIA LTDA", "CERT ALFA")


def criar_planilha(caminho, linhas=(ALFA, BETA, GAMA), status=None, cabecalho=True,
                   aba="Empresas"):
    """Gera um .xlsx com a aba de empresas.

    `status` mapeia indice da linha (0-based) -> (valor_D, valor_E).
    """
    status = status or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba
    if cabecalho:
        ws.append(CABECALHO_EMPRESAS)
    for i, (cnpj, empresa, cert) in enumerate(linhas):
        d, e = status.get(i, ("", ""))
        ws.append([cnpj, empresa, cert, d, e])
    wb.save(caminho)
    wb.close()
    return caminho


def linhas_de_debito(cnpj, quantidade=2):
    return [
        {
            "cnpj": cnpj, "tipo": f"TIPO {n}", "tributo": f"TRIB {n}",
            "receita": f"{1000 + n}", "pa_ex": f"0{n}/2026", "dt_vcto": f"1{n}/01/2026",
            "valor_original": f"{n}00,00", "saldo": f"{n}50,00",
        }
        for n in range(1, quantidade + 1)
    ]


def linhas_de_processo(cnpj, quantidade=2):
    return [
        {
            "cnpj": cnpj, "tipo": f"TIPO {n}", "receita": f"{2000 + n}",
            "pa_ex": f"0{n}/2026", "dt_vcto": f"2{n}/01/2026",
            "valor_original": f"{n}00,00", "saldo": f"{n}50,00",
            "processo_credito": f"PROC-{n}",
        }
        for n in range(1, quantidade + 1)
    ]


def ler_aba(caminho, nome):
    wb = openpyxl.load_workbook(caminho)
    try:
        if nome not in wb.sheetnames:
            return None
        return [list(linha) for linha in wb[nome].iter_rows(values_only=True)]
    finally:
        wb.close()
