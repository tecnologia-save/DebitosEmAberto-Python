"""A SEMANTICA das oito colunas de cada aba de detalhe.

Escrito ANTES de qualquer mudanca da fatia 14A.1 e commitado antes dela.

A 14A fechou a aba 'Empresas' e deixou um residual: das oito colunas de
'Processos Fiscais', so seis sao conferidas. As duas de fora divergiam entre o
codigo e o `PLANILHA MODELO.xlsx`, e eu me recusei a asserir o que as fontes
nao sustentavam.

Esta fatia rastreia cada valor ate quem o produz. Nao ha voto entre artefatos:
quem decide o que uma coluna significa e o dado que cai nela.

Nenhum teste abre navegador, contata o portal ou usa dado de cliente. As
sentinelas nomeiam o proprio campo, para que a coluna onde elas caem seja a
resposta.
"""
import pathlib

import openpyxl
import pytest
from planilhas_sinteticas import ALFA, ler_aba

from automation import consulta_fiscal, planilha

RAIZ = pathlib.Path(__file__).resolve().parent.parent
MODELO = RAIZ / "PLANILHA MODELO.xlsx"
CERT_ALFA = "CERT ALFA"

CABECALHO_EMPRESAS = ["CNPJ", "EMPRESA", "CERTIFICADO", "DÉBITOS",
                      "PROCESSOS FISCAIS"]


def montar(caminho, extras=()):
    """Uma planilha valida, com as abas de detalhe que o teste pedir."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(CABECALHO_EMPRESAS)
    ws.append([ALFA[0], ALFA[1], CERT_ALFA, "", ""])
    for nome, cabecalho, corpo in extras:
        outra = wb.create_sheet(nome)
        outra.append(list(cabecalho))
        for linha in corpo:
            outra.append(list(linha))
    wb.save(caminho)
    wb.close()
    return str(caminho)


def cabecalho_do_modelo(nome):
    wb = openpyxl.load_workbook(MODELO)
    try:
        return list(next(wb[nome].iter_rows(max_row=1, values_only=True)))
    finally:
        wb.close()


@pytest.fixture
def sessao():
    s = planilha.SessaoPlanilha()
    yield s
    s.descartar()


# Cada valor diz o nome do proprio campo. Onde ele cair, e o que aquela coluna
# significa na pratica.
def debito_marcado(cnpj=ALFA[0]):
    return {campo: f"<{campo}>" for campo in planilha.CAMPOS_DEBITO} | {"cnpj": cnpj}


def processo_marcado(cnpj=ALFA[0]):
    return {campo: f"<{campo}>" for campo in planilha.CAMPOS_PROCESSO} | {"cnpj": cnpj}


# ── §2 · o mapa das oito colunas de 'Débitos' ────────────────────────────────

def test_2_a_ordem_dos_campos_de_debito_e_a_do_cabecalho(tmp_path, sessao):
    """Posicao a posicao: qual CAMPO cai sob qual CABECALHO.

    Nao e teste de comprimento: cada valor carrega o nome do proprio campo.
    """
    caminho = montar(tmp_path / "p.xlsx")
    sessao.abrir(caminho)

    sessao.anexar_debitos([debito_marcado()])
    sessao.gravar()

    linhas = ler_aba(caminho, "Débitos")
    sob = dict(zip(linhas[0], linhas[1], strict=True))

    assert sob == {
        "CNPJ": ALFA[0],
        "TIPO": "<tipo>",
        "TRIBUTO": "<tributo>",
        "Rec.": "<receita>",
        "PA/Ex.": "<pa_ex>",
        "Dt.Vcto.": "<dt_vcto>",
        "Valor Original": "<valor_original>",
        "Saldo Devedor": "<saldo>",
    }


def test_2_e_o_produtor_de_cada_campo_de_debito_esta_no_extrator():
    """A origem, no JS que le a tabela do DCTFWeb. Nao ha aqui uma inferencia
    pelo nome da variavel: os rotulos estao no proprio portal."""
    fonte = consulta_fiscal.__file__
    codigo = pathlib.Path(fonte).read_text(encoding="utf-8")
    extrator = codigo[codigo.index("def _linhas_da_tabela_dctfweb"):]
    extrator = extrator[: extrator.index("\ndef ")]

    assert "if (labelText === 'Tributo')              tributo" in extrator
    assert "if (labelText === 'Valor original (R$)')  valor_original" in extrator
    assert "const tdReceita = tds.find(td => td.classList.contains('text-nowrap'))" \
        in extrator
    assert "const pa_ex   = tdsSimples[0]" in extrator
    assert "const dt_vcto = tdsSimples[1]" in extrator
    assert "cnpj, tipo, tributo, receita," in extrator


def test_2_e_o_modelo_concorda_nas_oito_primeiras_de_debitos():
    """As tres fontes — cabecalho do codigo, ordem dos campos e modelo — dizem a
    mesma coisa. A nona coluna do modelo o codigo nunca escreve."""
    modelo = cabecalho_do_modelo("Débitos")

    assert modelo[:8] == planilha.CABECALHO_DEBITOS
    assert len(modelo) == 9
    assert modelo[8] == "Informações Complementares"


# ── §3 · o mapa das oito colunas de 'Processos Fiscais' ──────────────────────

def test_3_a_ordem_dos_campos_de_processo_e_a_do_cabecalho(tmp_path, sessao):
    """A mesma prova, e e nela que a oitava posicao aparece: `processo_credito`
    cai sob o cabecalho que o CODIGO escreve."""
    caminho = montar(tmp_path / "p.xlsx")
    sessao.abrir(caminho)

    sessao.anexar_processos([processo_marcado()])
    sessao.gravar()

    linhas = ler_aba(caminho, "Processos Fiscais")
    sob = dict(zip(linhas[0], linhas[1], strict=True))

    assert sob == {
        "CNPJ": ALFA[0],
        "TIPO": "<tipo>",
        "RECEITA": "<receita>",
        "PA/Ex.": "<pa_ex>",
        "Dt.Vcto.": "<dt_vcto>",
        "Valor Original": "<valor_original>",
        "Saldo Devedor": "<saldo>",
        "Processo de Crédito": "<processo_credito>",
    }


# ── §4 · o rastro do OITAVO valor ────────────────────────────────────────────

def test_4_o_oitavo_valor_vem_do_PROCESSO_DE_CREDITO_do_portal():
    """A prova da semantica, no produtor.

    O valor sai de um elemento que o portal so revela ao clicar num botao cujo
    `aria-label` e 'Expandir processo de crédito', e o elemento revelado e
    `div.processo-credito`. Nao ha ambiguidade possivel: nem o botao, nem o
    div, nem a variavel falam de informacao complementar.
    """
    codigo = pathlib.Path(consulta_fiscal.__file__).read_text(encoding="utf-8")
    card = codigo[codigo.index("def _linhas_do_card"):]

    assert 'aria-label="Expandir processo de crédito"' in card
    assert "page.locator('div.processo-credito')" in card
    assert "processo_credito = (div_proc.text_content() or \"\").strip()" in card


def test_4_e_ele_atravessa_ate_a_OITAVA_posicao():
    """Do extrator ate a celula, sem trocar de nome no caminho."""
    codigo = pathlib.Path(consulta_fiscal.__file__).read_text(encoding="utf-8")

    assert "_linhas_da_tabela_do_card(page, cnpj, processo_credito)" in codigo
    assert "valor_original, saldo, processo_credito," in codigo

    assert planilha.CAMPOS_PROCESSO[7] == "processo_credito"
    assert planilha.CABECALHO_PROCESSOS[7] == "Processo de Crédito"


def test_4_informacoes_complementares_e_o_nome_do_CARD_e_nao_de_uma_coluna():
    """E aqui esta, provavelmente, a origem do erro do modelo.

    A expressao EXISTE no portal — mas nomeia o botao que ABRE o card do
    processo fiscal, aquele que lista os debitos dele. Nao nomeia campo nenhum
    dentro da tabela.

    Sao dois botoes diferentes, em dois momentos diferentes:

        'Expandir informações complementares do processo fiscal'  -> abre o card
        'Expandir processo de crédito'                            -> da o valor

    O valor da oitava coluna sai do segundo. Quem copiou o rotulo do primeiro
    para o cabecalho pegou o nome da tela, e nao o do dado.
    """
    codigo = pathlib.Path(consulta_fiscal.__file__).read_text(encoding="utf-8")

    assert codigo.count("informações complementares") == 1
    rotulo_do_card = ("Expandir informações complementares do "
                      "processo fiscal")
    assert f'aria-label="{rotulo_do_card}"' in codigo
    assert "CARD_PROCESSO" in codigo

    # E nenhum CAMPO se chama assim: os oito de cada aba estao nomeados, e
    # nenhum deles e informacao complementar.
    assert "informac" not in "".join(planilha.CAMPOS_PROCESSO)
    assert "informac" not in "".join(planilha.CAMPOS_DEBITO)


# ── §7 · a coluna 5, e a inconsistencia do proprio modelo ────────────────────

def test_7_a_coluna_5_do_modelo_foi_CANONIZADA():
    """ANTES o modelo trazia 'Dt.Vcto.' na aba de Débitos e 'Dt. Vcto' na de
    Processos Fiscais — discordava de SI MESMO, e era isso que tornava a segunda
    forma um erro de digitacao, e nao um contrato.

    A forma canonica e a do codigo, e agora as duas abas do modelo a usam.
    """
    debitos = cabecalho_do_modelo("Débitos")
    processos = cabecalho_do_modelo("Processos Fiscais")

    assert planilha.CABECALHO_DEBITOS[5] == "Dt.Vcto."
    assert planilha.CABECALHO_PROCESSOS[4] == "Dt.Vcto."

    assert debitos[5] == "Dt.Vcto."
    assert processos[4] == "Dt.Vcto."


def test_7_e_a_grafia_ANTIGA_continua_aceita_como_alias(tmp_path, sessao):
    """§7, opcao A: quem ja tem planilha feita a partir do modelo antigo nao e
    recusado por causa de um ponto final. E um conjunto FINITO de duas strings
    comprovadas — nao ha regex, nem strip de pontuacao, nem aproximacao."""
    antiga = [*planilha.CABECALHO_PROCESSOS]
    antiga[4] = "Dt. Vcto"
    caminho = montar(tmp_path / "p.xlsx",
                     extras=[("Processos Fiscais", antiga, [])])

    sessao.abrir(caminho)

    assert sessao.wb is not None


def test_7_e_o_alias_e_uma_LISTA_e_nao_uma_regra(tmp_path, sessao):
    """Uma terceira grafia qualquer nao passa: o conjunto e fechado."""
    inventada = [*planilha.CABECALHO_PROCESSOS]
    inventada[4] = "Dt Vencimento"
    caminho = montar(tmp_path / "p.xlsx",
                     extras=[("Processos Fiscais", inventada, [])])

    with pytest.raises(planilha.PlanilhaIndisponivel):
        sessao.abrir(caminho)


def test_7_e_a_coluna_7_diverge_so_por_ESPACO():
    """Essa a 14A ja resolve: espaco e normalizado."""
    processos = cabecalho_do_modelo("Processos Fiscais")

    assert processos[6] == "Saldo Devedor "
    assert processos[6].strip() == planilha.CABECALHO_PROCESSOS[6]


# ── §10 · o residual que a 14A deixou aberto ─────────────────────────────────

def test_10_seis_colunas_certas_NAO_bastam_mais(tmp_path, sessao):
    """ANTES a oitava coluna podia dizer outra coisa: a validacao nao olhava
    para ela, e `processo_credito` era gravado sob o cabecalho errado.

    §11: nada e appendado, nada e alterado, nada e salvo.
    """
    cabecalho_alheio = [*planilha.CABECALHO_PROCESSOS[:7], "OUTRA COISA"]
    caminho = montar(tmp_path / "p.xlsx",
                     extras=[("Processos Fiscais", cabecalho_alheio, [])])
    antes = ler_aba(caminho, "Processos Fiscais")

    with pytest.raises(planilha.PlanilhaIndisponivel):
        sessao.abrir(caminho)

    assert ler_aba(caminho, "Processos Fiscais") == antes
    assert sessao.wb is None


def test_10_e_o_MODELO_deixou_de_ter_esse_problema(tmp_path, sessao):
    """ANTES nao era um caso hipotetico: era o template que o projeto
    distribui. Agora a aba dele recebe cada valor sob a coluna certa.

    §9: copia sintetica do cabecalho do modelo, sem dado nenhum.
    """
    caminho = montar(
        tmp_path / "p.xlsx",
        extras=[("Processos Fiscais", cabecalho_do_modelo("Processos Fiscais"),
                 [])],
    )
    sessao.abrir(caminho)

    sessao.anexar_processos([processo_marcado()])
    sessao.gravar()

    linhas = ler_aba(caminho, "Processos Fiscais")
    sob = dict(zip(linhas[0], linhas[1], strict=True))

    assert sob["Processo de Crédito"] == "<processo_credito>"
    assert sob["Dt.Vcto."] == "<dt_vcto>"


# ── §13 · o README, como ele esta ────────────────────────────────────────────

def test_13_a_planilha_QUE_O_README_MANDA_MONTAR_e_aceita(tmp_path, sessao):
    """ANTES o README descrevia quatro colunas, com a D chamada de 'RESULTADO' —
    uma planilha que o validador recusa.

    §15: marcador estrutural minimo, e nao busca de prosa. O teste le a TABELA
    do README, monta a planilha que ela descreve, e exige que ela passe. Se a
    documentacao voltar a divergir do produto, e aqui que aparece.
    """
    texto = (RAIZ / "README.md").read_text(encoding="utf-8")
    tabela = texto[texto.index("## Formato da Planilha"):]
    tabela = tabela[: tabela.index("## Execução")]

    colunas = []
    for linha in tabela.splitlines():
        partes = [pedaco.strip() for pedaco in linha.split("|")]
        if len(partes) > 3 and len(partes[1]) == 1 and partes[1].isalpha():
            colunas.append(partes[2])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(colunas)
    ws.append([ALFA[0], ALFA[1], CERT_ALFA, "", ""][: len(colunas)])
    caminho = str(tmp_path / "readme.xlsx")
    wb.save(caminho)
    wb.close()

    sessao.abrir(caminho)

    assert sessao.wb is not None


# ── §8 · §12 · o que a fatia passou a garantir ───────────────────────────────

def test_8_a_aba_CRIADA_do_zero_e_semanticamente_coerente(tmp_path, sessao):
    """§8: criar, appendar, e cada valor sob o cabecalho correspondente. As duas
    abas, e nao so a que tinha o problema."""
    caminho = montar(tmp_path / "p.xlsx")
    sessao.abrir(caminho)

    sessao.anexar_debitos([debito_marcado()])
    sessao.anexar_processos([processo_marcado()])
    sessao.gravar()
    sessao.descartar()

    for aba, campos in (("Débitos", planilha.CAMPOS_DEBITO),
                        ("Processos Fiscais", planilha.CAMPOS_PROCESSO)):
        linhas = ler_aba(caminho, aba)
        valores = list(linhas[1])
        assert valores[0] == ALFA[0]
        assert valores[1:] == [f"<{campo}>" for campo in campos[1:]]

    # E o que foi escrito e reaberto sem reclamacao.
    planilha.validar_recurso(caminho)


def test_12_as_oito_de_DEBITOS_tambem_passaram_a_ser_conferidas(tmp_path,
                                                                 sessao):
    """§12: o metodo foi aplicado as duas abas. Débitos nao tinha divergencia —
    e agora tambem nao tem coluna sem conferencia."""
    cabecalho_alheio = [*planilha.CABECALHO_DEBITOS[:7], "OUTRA COISA"]
    caminho = montar(tmp_path / "p.xlsx",
                     extras=[("Débitos", cabecalho_alheio, [])])

    with pytest.raises(planilha.PlanilhaIndisponivel):
        sessao.abrir(caminho)


def test_12_e_a_NONA_coluna_de_debitos_continua_livre(tmp_path, sessao):
    """§17: nao se confere coluna que a automacao nao escreve. O modelo traz uma
    nona em `Débitos`, e ela e assunto de quem monta a planilha."""
    com_nona = [*planilha.CABECALHO_DEBITOS, "QUALQUER COISA MINHA"]
    caminho = montar(tmp_path / "p.xlsx",
                     extras=[("Débitos", com_nona, [])])

    sessao.abrir(caminho)

    assert sessao.wb is not None


def test_17_toda_coluna_ESCRITA_tem_cabecalho_conferido():
    """O criterio do §17, em uma frase: a automacao escreve oito valores em cada
    aba, e as oito posicoes sao conferidas."""
    assert len(planilha.CAMPOS_DEBITO) == len(planilha.CABECALHO_DEBITOS) == 8
    assert len(planilha.CAMPOS_PROCESSO) == len(planilha.CABECALHO_PROCESSOS) == 8

    fonte = (RAIZ / "automation" / "planilha.py").read_text(encoding="utf-8")
    assert "dict(enumerate(cabecalho))" in fonte, "todas as posicoes, e nao um subconjunto"
