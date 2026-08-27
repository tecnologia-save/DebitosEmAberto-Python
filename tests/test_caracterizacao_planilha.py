"""Caracterizacao da integracao de planilha COMO ELA E HOJE.

Escrito ANTES de extrair, depois da microcorrecao do F821. Usa .xlsx reais
gerados em tmp_path — porque o que se prova aqui e o contrato de PERSISTENCIA, e
um workbook falso provaria outra coisa.

Nenhuma planilha de cliente. Todos os CNPJs, empresas e certificados sao
ficticios.

A pergunta central: se a execucao cair no meio, o que sobrevive, e como a
proxima execucao sabe onde retomar?
"""
import openpyxl
import pytest
from planilhas_sinteticas import (
    ALFA,
    BETA,
    GAMA,
    criar_planilha,
    ler_aba,
    linhas_de_debito,
    linhas_de_processo,
)

from automation import planilha as planilha_mod

LIMPA = {"caminho": None, "wb": None, "sujo": False, "status": None}

# ── CHARACTERIZATION_TARGET_CHANGE (fatias 9B) ───────────────────────────────
# Primeiro os wrappers de coluna do `main` sairam; agora saiu tambem a SESSAO de
# planilha do `main` — o dono dela e o app. Estes testes caracterizam a
# integracao de planilha, entao passam a fala-la diretamente. Nenhuma afirmacao
# mudou: as mesmas celulas, as mesmas abas, o mesmo arquivo no disco.

_SESSAO = planilha_mod.SessaoPlanilha()


def _wb_sessao(caminho):
    if _SESSAO.precisa_abrir(caminho):
        fechar_planilha()
        _SESSAO.abrir(caminho)
    return _SESSAO.wb


def salvar_planilha():
    """O adapter de gravacao do `main`, agora local ao teste que o caracteriza."""
    if not _SESSAO.precisa_gravar():
        return False
    try:
        _SESSAO.gravar()
    except Exception as e:                                        # noqa: BLE001
        print(f"    [!] Falha ao salvar a planilha: {type(e).__name__}: {e}")
        return False
    _SESSAO.marcar_gravado()
    return True


def fechar_planilha():
    salvar_planilha()
    _SESSAO.descartar()


def mapa_status(caminho):
    _wb_sessao(caminho)
    return _SESSAO.mapa_status(caminho)


def ler_e_ordenar(caminho):
    return planilha_mod.ler_e_ordenar(caminho)[0]


def filtrar_pendentes(df, caminho):
    from automation.status_portal import status_encerra_linha

    return planilha_mod.linhas_pendentes(df, mapa_status(caminho), status_encerra_linha)


def _relatar_status(cnpj, rotulo, valor, gravou):
    if gravou:
        print(f"    [✓] Coluna {rotulo} → '{valor}'  (CNPJ {cnpj})")
        return
    print(f"    [!] CNPJ {cnpj} não encontrado na planilha para escrita em {rotulo}.")


def _registrar(caminho, cnpj, rotulo, valor, metodo):
    _wb_sessao(caminho)
    _relatar_status(cnpj, rotulo, valor, getattr(_SESSAO, metodo)(cnpj))

_METODO_D = {
    planilha_mod.STATUS_CONCLUIDO: "registrar_debitos_concluidos",
    planilha_mod.STATUS_SEM_DEBITOS: "registrar_sem_debitos",
    planilha_mod.STATUS_DEBITOS_NAO_COMPENSAVEIS: "registrar_debitos_nao_compensaveis",
}


def escrever_coluna_d(caminho, cnpj, valor):
    _registrar(caminho, cnpj, "D", valor, _METODO_D[valor])


def escrever_coluna_e(caminho, cnpj, valor):
    assert valor == planilha_mod.STATUS_SEM_PROCESSOS
    _registrar(caminho, cnpj, "E", valor, "registrar_sem_processos")


def escrever_aba_debitos(caminho, dados):
    _wb_sessao(caminho)
    _SESSAO.anexar_debitos(dados)


def escrever_aba_processos_fiscais(caminho, dados):
    _wb_sessao(caminho)
    _SESSAO.anexar_processos(dados)


def ler_status_cnpj(caminho, cnpj):
    return mapa_status(caminho).get(cnpj, ("", ""))


@pytest.fixture(autouse=True)
def sessao_limpa():
    """A sessao e um dict global de modulo — cada teste comeca e termina limpo."""
    _SESSAO.estado.update(LIMPA)
    yield
    try:
        fechar_planilha()
    except (OSError, KeyError, ValueError):
        # Um teste pode deixar a sessao apontando para algo que nao grava; a
        # limpeza nao pode contaminar o proximo teste por causa disso.
        pass
    _SESSAO.estado.update(LIMPA)


@pytest.fixture
def planilha(tmp_path):
    return str(criar_planilha(tmp_path / "base.xlsx"))


# ── A · B · abertura e leitura ────────────────────────────────────────────────

def test_a_workbook_minimo_valido(planilha):
    df = ler_e_ordenar(planilha)

    assert len(df) == 3
    # Fatia 14A: o cabecalho sintetico passou a ser o do `PLANILHA MODELO.xlsx`.
    # Antes era invencao nossa, e a validacao de schema o recusaria.
    assert list(df.columns) == ["CNPJ", "EMPRESA", "CERTIFICADO", "DÉBITOS",
                                "PROCESSOS FISCAIS"]


def test_b_a_aba_empresas_e_lida_por_nome(planilha):
    """pd.read_excel com sheet_name='Empresas' — o nome importa, a posicao nao."""
    df = ler_e_ordenar(planilha)
    assert set(df["CNPJ"]) == {ALFA[0], BETA[0], GAMA[0]}


# ── C · aba ausente ───────────────────────────────────────────────────────────

def test_c_aba_empresas_ausente_na_leitura_pandas(tmp_path):
    caminho = str(criar_planilha(tmp_path / "sem_aba.xlsx", aba="Outra"))

    with pytest.raises(ValueError, match="Empresas"):
        ler_e_ordenar(caminho)


def test_c_aba_empresas_ausente_no_acesso_openpyxl(tmp_path):
    """ANTES: `openpyxl` levantava `KeyError` la dentro, ao procurar a aba —
    depois de o workbook ja estar aberto e adotado pela sessao.

    Fatia 14A: quem recusa agora e a abertura, com o erro de ENTRADA do
    projeto. O `KeyError` cru nunca mais chega a quem chamou.
    """
    caminho = str(criar_planilha(tmp_path / "sem_aba.xlsx", aba="Outra"))

    with pytest.raises(planilha_mod.PlanilhaIndisponivel):
        mapa_status(caminho)


def test_c_arquivo_inexistente(tmp_path):
    with pytest.raises(FileNotFoundError):
        mapa_status(str(tmp_path / "nao_existe.xlsx"))


# ── D · colunas ───────────────────────────────────────────────────────────────

def test_d_as_colunas_sao_posicionais_e_nao_nomeadas(tmp_path):
    """PLANILHA_POSSIBLE_DEFECT: A=CNPJ, C=certificado, D e E=status vem por
    POSICAO. Uma coluna a mais no inicio desloca tudo em silencio."""
    caminho = tmp_path / "deslocada.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(["ORDEM", "CNPJ", "EMPRESA", "CERTIFICADO", "DCTFWEB", "PROCESSOS"])
    ws.append([1, ALFA[0], ALFA[1], ALFA[2], "", ""])
    wb.save(caminho)
    wb.close()

    df = ler_e_ordenar(str(caminho))
    assert df.columns[0] == "ORDEM", "o codigo trataria isto como a coluna do CNPJ"
    assert df.columns[2] == "EMPRESA", "e isto como a coluna do certificado"


def test_d_menos_de_cinco_colunas_agora_e_RECUSADO(tmp_path):
    """ANTES: `mapa_status` protegia o acesso com `len(linha) > 3` e devolvia
    `("", "")` — a planilha sem as colunas de status era tratada como uma
    planilha inteiramente pendente, e a automacao escrevia em colunas que nao
    existiam no cabecalho.

    Fatia 14A: falta coluna, falta formato. PLANILHA_SCHEMA_FAIL_CLOSED.

    A protecao `len(linha) > 3` continua no codigo: ela vale para a LINHA, que
    pode ser mais curta que o cabecalho numa planilha valida.
    """
    caminho = tmp_path / "curta.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(["CNPJ", "EMPRESA", "CERTIFICADO"])
    ws.append([ALFA[0], ALFA[1], ALFA[2]])
    wb.save(caminho)
    wb.close()

    with pytest.raises(planilha_mod.PlanilhaIndisponivel):
        mapa_status(str(caminho))


def test_d_e_uma_LINHA_mais_curta_que_o_cabecalho_continua_valendo(tmp_path):
    """O outro lado, e o que a recusa acima nao pode ter levado junto: o
    cabecalho tem as cinco colunas e a linha tem tres. Isso e planilha em
    branco, nao planilha incompativel."""
    caminho = tmp_path / "linha_curta.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(["CNPJ", "EMPRESA", "CERTIFICADO", "DÉBITOS", "PROCESSOS FISCAIS"])
    ws.append([ALFA[0], ALFA[1], ALFA[2]])
    wb.save(caminho)
    wb.close()

    assert mapa_status(str(caminho))[ALFA[0]] == ("", "")


# ── E · F · pendentes e concluidos ────────────────────────────────────────────

def test_e_todas_pendentes_quando_d_e_e_estao_vazias(planilha):
    df = ler_e_ordenar(planilha)
    pendentes, concluidas = filtrar_pendentes(df, planilha)

    assert len(pendentes) == 3
    assert concluidas == 0


def test_f_linha_com_d_e_e_preenchidas_e_concluida(tmp_path):
    caminho = str(criar_planilha(
        tmp_path / "meio.xlsx", status={0: ("Concluído", "Sem Processos")}
    ))
    df = ler_e_ordenar(caminho)
    pendentes, concluidas = filtrar_pendentes(df, caminho)

    assert concluidas == 1
    assert ALFA[0] not in set(pendentes["CNPJ"])


def test_f_so_uma_das_duas_colunas_nao_conclui(tmp_path):
    """D preenchida e E vazia continua pendente — fara so Processos Fiscais."""
    caminho = str(criar_planilha(tmp_path / "parcial.xlsx", status={0: ("Concluído", "")}))
    df = ler_e_ordenar(caminho)
    pendentes, concluidas = filtrar_pendentes(df, caminho)

    assert concluidas == 0
    assert ALFA[0] in set(pendentes["CNPJ"])


def test_f_status_terminal_em_d_conclui_sozinho(tmp_path):
    """A recusa de procuracao encerra a linha sem depender da coluna E."""
    caminho = str(criar_planilha(
        tmp_path / "recusada.xlsx", status={0: ("Procuração sem autorização", "")}
    ))
    df = ler_e_ordenar(caminho)
    _, concluidas = filtrar_pendentes(df, caminho)

    assert concluidas == 1


# ── G · ordenacao ─────────────────────────────────────────────────────────────

def test_g_ordena_pelo_certificado_da_coluna_c(planilha):
    """Agrupar por certificado e o que permite um login servir varios CNPJs."""
    df = ler_e_ordenar(planilha)

    assert list(df["CERTIFICADO"]) == ["CERT ALFA", "CERT ALFA", "CERT BETA"]
    assert list(df["CNPJ"])[:2] == [ALFA[0], GAMA[0]], "estavel dentro do grupo"


# ── H · I · escrita das colunas D e E ─────────────────────────────────────────

def test_h_escrita_da_coluna_d(planilha):
    escrever_coluna_d(planilha, ALFA[0], "Concluído")
    salvar_planilha()

    assert ler_aba(planilha, "Empresas")[1][3] == "Concluído"


def test_i_escrita_da_coluna_e(planilha):
    escrever_coluna_e(planilha, BETA[0], "Sem Processos")
    salvar_planilha()

    assert ler_aba(planilha, "Empresas")[2][4] == "Sem Processos"


def test_h_a_escrita_sincroniza_o_mapa_em_memoria(planilha):
    mapa_status(planilha)
    escrever_coluna_d(planilha, ALFA[0], "Concluído")

    assert ler_status_cnpj(planilha, ALFA[0]) == ("Concluído", "")


def test_h_o_cnpj_e_localizado_com_qualquer_formatacao(tmp_path):
    caminho = str(criar_planilha(
        tmp_path / "formatado.xlsx", linhas=[("11.111.111/0001-91", "ALFA", "CERT")]
    ))
    escrever_coluna_d(caminho, ALFA[0], "Concluído")
    salvar_planilha()

    assert ler_aba(caminho, "Empresas")[1][3] == "Concluído"


def test_h_cnpj_ausente_nao_grava_e_apenas_avisa(planilha, capsys):
    """PLANILHA_POSSIBLE_DEFECT: o status e descartado em silencio funcional —
    so um print. Quem chamou nao tem como saber que nada foi gravado."""
    escrever_coluna_d(planilha, "99999999000199", "Concluído")

    assert "não encontrado na planilha" in capsys.readouterr().out
    assert _SESSAO.estado["sujo"] is False
    assert salvar_planilha() is False


# ── J · K · L · M · abas de detalhe ───────────────────────────────────────────

def test_j_a_aba_debitos_e_criada_com_cabecalho(planilha):
    escrever_aba_debitos(planilha, linhas_de_debito(ALFA[0]))
    salvar_planilha()

    aba = ler_aba(planilha, "Débitos")
    assert aba[0][:3] == ["CNPJ", "TIPO", "TRIBUTO"]
    assert len(aba) == 3


def test_k_append_preserva_o_que_ja_estava(planilha):
    escrever_aba_debitos(planilha, linhas_de_debito(ALFA[0]))
    escrever_aba_debitos(planilha, linhas_de_debito(BETA[0]))
    salvar_planilha()

    aba = ler_aba(planilha, "Débitos")
    assert len(aba) == 5, "cabecalho + 2 + 2"
    assert [linha[0] for linha in aba[1:]] == [ALFA[0]] * 2 + [BETA[0]] * 2


def test_l_a_aba_processos_fiscais_e_criada_com_cabecalho(planilha):
    escrever_aba_processos_fiscais(planilha, linhas_de_processo(ALFA[0]))
    salvar_planilha()

    aba = ler_aba(planilha, "Processos Fiscais")
    assert aba[0][-1] == "Processo de Crédito"
    assert len(aba) == 3


def test_m_append_em_processos_fiscais(planilha):
    escrever_aba_processos_fiscais(planilha, linhas_de_processo(ALFA[0], 1))
    escrever_aba_processos_fiscais(planilha, linhas_de_processo(BETA[0], 3))
    salvar_planilha()

    assert len(ler_aba(planilha, "Processos Fiscais")) == 5


def test_k_lacunas_no_meio_sao_preenchidas_antes_do_fim(planilha):
    """Documentado no codigo, mas surpreendente: linhas apagadas no meio recebem
    dados novos. As linhas de um CNPJ podem acabar espalhadas."""
    escrever_aba_debitos(planilha, linhas_de_debito(ALFA[0], 3))
    salvar_planilha()

    wb = openpyxl.load_workbook(planilha)
    for celula in wb["Débitos"][3]:
        celula.value = None
    wb.save(planilha)
    wb.close()
    fechar_planilha()

    escrever_aba_debitos(planilha, linhas_de_debito(BETA[0], 1))
    salvar_planilha()

    aba = ler_aba(planilha, "Débitos")
    assert aba[2][0] == BETA[0], "entrou na lacuna, nao no fim"
    assert len(aba) == 4, "nenhuma linha nova foi criada"


def test_k_linha_ocupada_nunca_e_sobrescrita(planilha):
    escrever_aba_debitos(planilha, linhas_de_debito(ALFA[0], 2))
    escrever_aba_debitos(planilha, linhas_de_debito(BETA[0], 2))
    salvar_planilha()

    aba = ler_aba(planilha, "Débitos")
    assert len([linha for linha in aba[1:] if linha[0]]) == 4


# ── N · O · salvamento e fechamento ───────────────────────────────────────────

def test_n_salvar_so_grava_quando_ha_alteracao(planilha):
    mapa_status(planilha)
    assert salvar_planilha() is False, "leitura nao suja a sessao"

    escrever_coluna_d(planilha, ALFA[0], "Concluído")
    assert salvar_planilha() is True
    assert salvar_planilha() is False, "ja gravou"


def test_n_o_mesmo_arquivo_e_sobrescrito(tmp_path):
    """Sem copia temporaria e sem escrita atomica: wb.save(caminho_recebido)."""
    caminho = str(criar_planilha(tmp_path / "base.xlsx"))
    escrever_coluna_d(caminho, ALFA[0], "Concluído")
    salvar_planilha()

    assert [p.name for p in tmp_path.iterdir()] == ["base.xlsx"]


def test_o_fechar_salva_o_que_estava_pendente(planilha):
    escrever_coluna_d(planilha, ALFA[0], "Concluído")
    fechar_planilha()

    assert ler_aba(planilha, "Empresas")[1][3] == "Concluído"
    assert _SESSAO.estado["wb"] is None
    assert _SESSAO.estado["status"] is None


def test_o_trocar_de_planilha_fecha_a_anterior(tmp_path):
    primeira = str(criar_planilha(tmp_path / "a.xlsx"))
    segunda = str(criar_planilha(tmp_path / "b.xlsx"))

    escrever_coluna_d(primeira, ALFA[0], "Concluído")
    mapa_status(segunda)

    assert ler_aba(primeira, "Empresas")[1][3] == "Concluído", "salvou ao trocar"
    assert _SESSAO.estado["caminho"] == segunda


# ── P · falha de salvamento ───────────────────────────────────────────────────

def test_p_falha_ao_salvar_mantem_a_sessao_suja(planilha, monkeypatch):
    escrever_coluna_d(planilha, ALFA[0], "Concluído")

    def falhar(*a, **k):
        raise PermissionError("arquivo aberto no Excel")

    monkeypatch.setattr(_SESSAO.estado["wb"], "save", falhar)

    assert salvar_planilha() is False
    assert _SESSAO.estado["sujo"] is True, "o dado nao foi descartado"


# ── Q · R · retomada e reexecucao ─────────────────────────────────────────────

def test_q_progresso_parcial_sobrevive_e_orienta_a_retomada(tmp_path):
    """RESUMABILITY_CONTRACT — o teste mais importante desta fatia.

    Simula uma execucao que processou ALFA por inteiro, GAMA pela metade, e caiu
    antes de BETA. A proxima execucao tem de: pular ALFA, refazer so o que falta
    de GAMA, e fazer BETA inteiro.
    """
    caminho = str(criar_planilha(tmp_path / "base.xlsx"))

    escrever_aba_debitos(caminho, linhas_de_debito(ALFA[0], 2))
    escrever_coluna_d(caminho, ALFA[0], "Concluído")
    escrever_coluna_e(caminho, ALFA[0], "Sem Processos")
    escrever_coluna_d(caminho, GAMA[0], "Concluído")
    fechar_planilha()   # queda simulada apos o fechamento do CNPJ

    df = ler_e_ordenar(caminho)
    pendentes, concluidas = filtrar_pendentes(df, caminho)

    assert concluidas == 1, "ALFA nao volta"
    assert set(pendentes["CNPJ"]) == {GAMA[0], BETA[0]}
    assert ler_status_cnpj(caminho, GAMA[0]) == ("Concluído", ""), "GAMA so Processos"
    assert len(ler_aba(caminho, "Débitos")) == 3, "os debitos de ALFA continuam la"


def test_r_reexecucao_completa_nao_reprocessa_nada(tmp_path):
    caminho = str(criar_planilha(
        tmp_path / "pronta.xlsx",
        status={i: ("Concluído", "Sem Processos") for i in range(3)},
    ))
    df = ler_e_ordenar(caminho)
    pendentes, concluidas = filtrar_pendentes(df, caminho)

    assert pendentes.empty
    assert concluidas == 3


def test_r_retomada_duplica_detalhe_se_o_status_nao_tiver_sido_gravado(tmp_path):
    """PLANILHA_POSSIBLE_DEFECT: as abas de detalhe sao append-only e nao tem
    deduplicacao por CNPJ.

    Se as linhas de detalhe forem gravadas e a coluna de status NAO — janela
    estreita, mas real — a retomada anexa tudo de novo. O contrato de retomada
    mora inteiro nas colunas D e E; as abas de detalhe nao participam dele.
    """
    caminho = str(criar_planilha(tmp_path / "base.xlsx"))

    escrever_aba_debitos(caminho, linhas_de_debito(ALFA[0], 2))
    fechar_planilha()   # detalhe no disco, coluna D em branco

    df = ler_e_ordenar(caminho)
    pendentes, _ = filtrar_pendentes(df, caminho)
    assert ALFA[0] in set(pendentes["CNPJ"]), "volta como pendente"

    escrever_aba_debitos(caminho, linhas_de_debito(ALFA[0], 2))
    salvar_planilha()

    assert len(ler_aba(caminho, "Débitos")) == 5, "as mesmas 2 linhas, duas vezes"


def test_r_cnpj_duplicado_na_aba_empresas(tmp_path):
    """PLANILHA_POSSIBLE_DEFECT: com o mesmo CNPJ em duas linhas, a ESCRITA vai
    para a primeira e a LEITURA do mapa reflete a ultima.

    Consequencia: a segunda linha nunca e marcada e volta pendente em toda
    execucao — retrabalho silencioso e sem fim.
    """
    caminho = str(criar_planilha(tmp_path / "dup.xlsx", linhas=[ALFA, ALFA]))

    escrever_coluna_d(caminho, ALFA[0], "Concluído")
    escrever_coluna_e(caminho, ALFA[0], "Sem Processos")
    fechar_planilha()

    empresas = ler_aba(caminho, "Empresas")
    assert empresas[1][3] == "Concluído", "primeira linha marcada"
    assert empresas[2][3] is None, "segunda linha intocada"

    df = ler_e_ordenar(caminho)
    _, concluidas = filtrar_pendentes(df, caminho)
    assert concluidas == 0, "o mapa reflete a ULTIMA linha, que esta vazia"
